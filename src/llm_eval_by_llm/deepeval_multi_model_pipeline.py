"""
DeepEval Multi-Model × Multi-Prompt Pipeline
==============================================
Runs multiple LLM models against source PDFs using various prompts from:
  1. The prompt library (zero_shot, few_shot, chain_of_thought, etc.)
  2. mCodeGPT-style prompts (BFOP, 2POP)

For each (model, prompt_version, case) combination:
  - Extracts text from source PDFs
  - Sends text + prompt to the LLM
  - Parses the JSON response for each element
  - Validates against source ground truth
  - Outputs a validation sheet matching the human validation format

Uses Confident AI DeepEval for evaluation metrics (hallucination, faithfulness).

Usage:
    python deepeval_multi_model_pipeline.py [--models gpt-4o,claude-3-sonnet] [--dry-run]
"""

import os
import sys
import json
import re
import time
import logging
import argparse
from pathlib import Path
from datetime import datetime
from typing import Optional
from collections import defaultdict

import numpy as np
import pandas as pd

# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------
PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_DIR = PROJECT_ROOT / "data" / "raw"
PROCESSED_DIR = PROJECT_ROOT / "data" / "processed"
PROMPT_LIBRARY_CSV = PROJECT_ROOT / "references" / "Prompts" / "prompt_library (1).csv"
PROMPT_LIBRARY_XLSX = PROCESSED_DIR / "prompt_library_updated_v5.xlsx"
OUTPUT_DIR = PROJECT_ROOT / "data" / "processed" / "deepeval_runs"
VALIDATION_RAW = RAW_DIR / "merged_llm_summary_validation_datasheet.xlsx"

# ---------------------------------------------------------------------------
# Logging
# ---------------------------------------------------------------------------
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-8s  %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Element definitions (must match validation datasheet columns)
# ---------------------------------------------------------------------------
ELEMENTS = [
    {"key": "lesion_size", "display": "Lesion Size", "source_col": "lesion_size_status_source"},
    {"key": "lesion_laterality", "display": "Lesion Laterality", "source_col": "laterality_status_source"},
    {"key": "lesion_location", "display": "Lesion Location", "source_col": "lesion_location_status_source"},
    {"key": "calcifications_asymmetry", "display": "Calcifications / Asymmetry", "source_col": "calcifications_asymmetry_status_source"},
    {"key": "additional_enhancement_mri", "display": "Additional Enhancement (MRI)", "source_col": "additional_enhancement_mri_status_source"},
    {"key": "extent", "display": "Extent", "source_col": "extent_status_source"},
    {"key": "accurate_clip_placement", "display": "Accurate Clip Placement", "source_col": "accurate_clip_placement_status_source"},
    {"key": "workup_recommendation", "display": "Workup Recommendation", "source_col": "workup_recommendation_status_source"},
    {"key": "lymph_node", "display": "Lymph Node", "source_col": "Lymph node_status_source"},
    {"key": "chronology_preserved", "display": "Chronology Preserved", "source_col": "chronology_preserved_status_source"},
    {"key": "biopsy_method", "display": "Biopsy Method", "source_col": "biopsy_method_status_source"},
    {"key": "invasive_component_size", "display": "Invasive Component Size (Pathology)", "source_col": "invasive_component_size_pathology_status_source"},
    {"key": "histologic_diagnosis", "display": "Histologic Diagnosis", "source_col": "histologic_diagnosis_status_source"},
    {"key": "receptor_status", "display": "Receptor Status", "source_col": "receptor_status_source"},
]

# Map from prompt library CSV element names → our element keys
CSV_ELEMENT_MAP = {
    "lesion_size": "lesion_size",
    "lesion_location": "lesion_location",
    "lesion_laterality": "lesion_laterality",
    "calcifications_asymmetry": "calcifications_asymmetry",
    "additional_enhancement_mri": "additional_enhancement_mri",
    "extent": "extent",
    "clip_placement": "accurate_clip_placement",
    "workup_recommendation": "workup_recommendation",
    "lymph_node": "lymph_node",
    "chronology": "chronology_preserved",
    "biopsy_method": "biopsy_method",
    "invasive_component_size": "invasive_component_size",
    "histologic_diagnosis": "histologic_diagnosis",
    "receptor_status": "receptor_status",
}

# Prompting approaches in the CSV
PROMPT_APPROACHES = [
    "zero_shot", "zero_shot2",
    "few_shot", "few_shot2",
    "chain_of_thought", "chain_of_thought2",
    "program_aided", "program_aided2",
    "rag", "rag2",
    "react", "react2",
]

# Surgeon name → folder mapping
SURGEON_NAME_TO_DIR = {
    "Barrio, Andrea": "Andrea Barrio",
    "Choi, Danny": "Danny Choi",
    "Capko, Deborah": "Deborah Capko",
    "Allen, Lisa": "Lisa Allen",
    "Goel, Neha": "Neha Goel",
    "Bayard, Solange": "Solange Bayard",
    "Downs-Canner, Stephanie": "Stephanie Downs-Canner",
    "Pawloski": "Pawloski",
}

# ---------------------------------------------------------------------------
# Supported models
# ---------------------------------------------------------------------------
DEFAULT_MODELS = ["gpt-4o", "gpt-4o-mini", "claude-3-5-sonnet-20241022"]


# ===================================================================
# PDF / DOCX text extraction (reused from Phase 1)
# ===================================================================
def extract_text_from_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        return "\n".join(p.extract_text() or "" for p in reader.pages).strip()
    except Exception as exc:
        log.warning("PDF read failed %s: %s", path.name, exc)
        return ""


def extract_text_from_docx(path: Path) -> str:
    try:
        import docx
        doc = docx.Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs).strip()
    except Exception as exc:
        log.warning("DOCX read failed %s: %s", path.name, exc)
        return ""


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(path)
    elif ext in (".docx", ".doc"):
        return extract_text_from_docx(path)
    return ""


def classify_report(filename: str) -> str:
    fn = filename.lower()
    if "hpi_ai" in fn:
        return "hpi_ai"
    if "hpi_human" in fn:
        return "hpi_human"
    if "path" in fn or "biopsy" in fn or "bopsy" in fn:
        return "pathology"
    if "imaging" in fn or "mammo" in fn or "mri" in fn or "us" in fn:
        return "radiology"
    return "other"


def get_source_text(case_dir: Path) -> str:
    """Extract and concatenate text from source documents (not hpi_ai/hpi_human)."""
    texts = []
    for f in sorted(case_dir.iterdir()):
        if f.name.startswith(".") or f.name.startswith("_"):
            continue
        if f.suffix.lower() not in (".pdf", ".docx", ".doc"):
            continue
        rtype = classify_report(f.name)
        if rtype in ("hpi_ai", "hpi_human"):
            continue
        text = extract_text(f)
        if text:
            texts.append(f"--- {f.name} ---\n{text}")
    return "\n\n".join(texts)


# ===================================================================
# Prompt loading
# ===================================================================
def load_prompt_library() -> pd.DataFrame:
    """Load the prompt library from CSV (fallback from xlsx if permission denied)."""
    if PROMPT_LIBRARY_CSV.exists():
        df = pd.read_csv(PROMPT_LIBRARY_CSV)
        log.info("Loaded prompt library CSV: %d elements × %d approaches", len(df), len(df.columns) - 1)
        return df
    try:
        df = pd.read_excel(PROMPT_LIBRARY_XLSX)
        log.info("Loaded prompt library XLSX: %d rows", len(df))
        return df
    except Exception as exc:
        log.error("Cannot load prompt library: %s", exc)
        return pd.DataFrame()


def get_prompts_for_approach(prompt_df: pd.DataFrame, approach: str) -> dict:
    """
    Return {element_key: prompt_text} for a given prompting approach column.
    """
    if approach not in prompt_df.columns:
        return {}
    prompts = {}
    for _, row in prompt_df.iterrows():
        csv_element = str(row.get("element", "")).strip()
        element_key = CSV_ELEMENT_MAP.get(csv_element)
        if element_key and pd.notna(row[approach]):
            prompts[element_key] = str(row[approach])
    return prompts


# ===================================================================
# mCodeGPT-style prompts: BFOP and 2POP
# ===================================================================
BFOP_SYSTEM_PROMPT = """You are a clinical data extraction assistant using the Breadth-First Ontology Pruner (BFOP) method.

BFOP Strategy:
1. Start with top-level clinical categories (Radiology findings, Pathology findings).
2. For each category, determine if relevant information EXISTS in the source text.
3. If a category is ABSENT, skip all its child elements entirely.
4. If PRESENT, drill into specific elements (size, location, laterality, etc.).
5. This hierarchical pruning ensures efficient and accurate extraction.

For each element, output one of:
- The extracted value (if clearly present in source text)
- "Not reported" (if the element is not mentioned)
- "Unclear (illegible)" (if mentioned but unreadable/ambiguous)

CRITICAL: Do NOT fabricate or infer information not explicitly stated in the source documents.
Output ONLY valid JSON."""

BFOP_ELEMENT_PROMPTS = {
    "lesion_size": 'BFOP Level 2 — Radiology > Lesion Measurements: Extract lesion size(s) in cm. Return {"lesion_size_cm":"..."}.',
    "lesion_laterality": 'BFOP Level 2 — Radiology > Laterality: Extract laterality (left/right/bilateral). Return {"laterality":"..."}.',
    "lesion_location": 'BFOP Level 2 — Radiology > Location: Extract location (quadrant, clock-face, depth). Return {"lesion_location":"..."}.',
    "calcifications_asymmetry": 'BFOP Level 2 — Radiology > Calcifications/Asymmetry: Extract calcification or asymmetry findings. Return {"calcifications_asymmetry":"..."}.',
    "additional_enhancement_mri": 'BFOP Level 2 — Radiology > MRI Enhancement: Extract additional MRI enhancement findings. Return {"additional_enhancement_mri":"..."}.',
    "extent": 'BFOP Level 2 — Radiology > Extent: Extract extent of disease. Return {"extent":"..."}.',
    "accurate_clip_placement": 'BFOP Level 2 — Radiology > Clip Placement: Extract clip placement details. Return {"clip_placement":"..."}.',
    "workup_recommendation": 'BFOP Level 2 — Radiology > Recommendation: Extract workup recommendations. Return {"workup_recommendation":"..."}.',
    "lymph_node": 'BFOP Level 2 — Radiology > Lymph Nodes: Extract lymph node findings. Return {"lymph_node":"..."}.',
    "chronology_preserved": 'BFOP Level 2 — Radiology > Chronology: Determine if chronological order of events is preserved. Return {"chronology_preserved":"yes/no/unclear"}.',
    "biopsy_method": 'BFOP Level 2 — Pathology > Biopsy Method: Extract biopsy method (core, stereotactic, FNA, etc.). Return {"biopsy_method":"..."}.',
    "invasive_component_size": 'BFOP Level 2 — Pathology > Invasive Size: Extract invasive component size from pathology. Return {"invasive_component_size_cm":"..."}.',
    "histologic_diagnosis": 'BFOP Level 2 — Pathology > Histology: Extract histologic diagnosis. Return {"histologic_diagnosis":"..."}.',
    "receptor_status": 'BFOP Level 2 — Pathology > Receptors: Extract receptor status (ER, PR, HER2, Ki-67). Return {"receptor_status":"..."}.',
}

TWO_POP_SYSTEM_PROMPT = """You are a clinical data extraction assistant using the Two-Phase Ontology Parser (2POP) method.

2POP Strategy:
Phase 1 — PRESENCE CHECK: For each clinical element, determine if it is present in the source text (yes/no).
Phase 2 — DETAIL EXTRACTION: For elements confirmed as present, extract the specific value.

This two-phase approach reduces hallucination by first confirming presence before attempting extraction.

For each element, output one of:
- The extracted value (if confirmed present and clearly stated)
- "Not reported" (if confirmed absent)
- "Unclear (illegible)" (if present but unreadable/ambiguous)

CRITICAL: Do NOT fabricate or infer information not explicitly stated in the source documents.
Output ONLY valid JSON."""

TWO_POP_ELEMENT_PROMPTS = {
    "lesion_size": '2POP Phase 1: Is lesion size mentioned? Phase 2: If yes, extract size in cm. Return {"present": true/false, "lesion_size_cm":"..."}.',
    "lesion_laterality": '2POP Phase 1: Is laterality mentioned? Phase 2: If yes, extract. Return {"present": true/false, "laterality":"..."}.',
    "lesion_location": '2POP Phase 1: Is lesion location mentioned? Phase 2: If yes, extract. Return {"present": true/false, "lesion_location":"..."}.',
    "calcifications_asymmetry": '2POP Phase 1: Are calcifications/asymmetry mentioned? Phase 2: If yes, extract. Return {"present": true/false, "calcifications_asymmetry":"..."}.',
    "additional_enhancement_mri": '2POP Phase 1: Is MRI enhancement mentioned? Phase 2: If yes, extract. Return {"present": true/false, "additional_enhancement_mri":"..."}.',
    "extent": '2POP Phase 1: Is extent of disease mentioned? Phase 2: If yes, extract. Return {"present": true/false, "extent":"..."}.',
    "accurate_clip_placement": '2POP Phase 1: Is clip placement mentioned? Phase 2: If yes, extract. Return {"present": true/false, "clip_placement":"..."}.',
    "workup_recommendation": '2POP Phase 1: Is a workup recommendation mentioned? Phase 2: If yes, extract. Return {"present": true/false, "workup_recommendation":"..."}.',
    "lymph_node": '2POP Phase 1: Are lymph node findings mentioned? Phase 2: If yes, extract. Return {"present": true/false, "lymph_node":"..."}.',
    "chronology_preserved": '2POP Phase 1: Is chronological ordering present? Phase 2: If yes, is it preserved? Return {"present": true/false, "chronology_preserved":"yes/no"}.',
    "biopsy_method": '2POP Phase 1: Is biopsy method mentioned? Phase 2: If yes, extract. Return {"present": true/false, "biopsy_method":"..."}.',
    "invasive_component_size": '2POP Phase 1: Is invasive component size mentioned? Phase 2: If yes, extract. Return {"present": true/false, "invasive_component_size_cm":"..."}.',
    "histologic_diagnosis": '2POP Phase 1: Is histologic diagnosis mentioned? Phase 2: If yes, extract. Return {"present": true/false, "histologic_diagnosis":"..."}.',
    "receptor_status": '2POP Phase 1: Is receptor status mentioned? Phase 2: If yes, extract. Return {"present": true/false, "receptor_status":"..."}.',
}


# ===================================================================
# LLM call abstraction
# ===================================================================
def call_llm(
    model_name: str,
    system_prompt: str,
    user_prompt: str,
    temperature: float = 0.0,
    max_tokens: int = 1024,
) -> str:
    """
    Call an LLM via the appropriate SDK. Supports OpenAI and Anthropic models.
    Returns the raw text response.
    """
    if model_name.startswith("gpt") or model_name.startswith("o1") or model_name.startswith("ft:"):
        return _call_openai(model_name, system_prompt, user_prompt, temperature, max_tokens)
    elif "claude" in model_name:
        return _call_anthropic(model_name, system_prompt, user_prompt, temperature, max_tokens)
    elif "mistral" in model_name.lower():
        return _call_mistral(model_name, system_prompt, user_prompt, temperature, max_tokens)
    else:
        log.warning("Unknown model %s, trying OpenAI-compatible API", model_name)
        return _call_openai(model_name, system_prompt, user_prompt, temperature, max_tokens)


def _call_openai(model, system_prompt, user_prompt, temperature, max_tokens) -> str:
    from openai import OpenAI
    client = OpenAI()
    response = client.chat.completions.create(
        model=model,
        temperature=temperature,
        max_tokens=max_tokens,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
    )
    return response.choices[0].message.content or ""


def _call_anthropic(model, system_prompt, user_prompt, temperature, max_tokens) -> str:
    from anthropic import Anthropic
    client = Anthropic()
    response = client.messages.create(
        model=model,
        max_tokens=max_tokens,
        temperature=temperature,
        system=system_prompt,
        messages=[{"role": "user", "content": user_prompt}],
    )
    return response.content[0].text if response.content else ""


def _call_mistral(model, system_prompt, user_prompt, temperature, max_tokens) -> str:
    from mistralai import Mistral
    client = Mistral(api_key=os.environ.get("MISTRAL_API_KEY", ""))
    response = client.chat.complete(
        model=model,
        temperature=temperature,
        max_tokens=max_tokens,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
    )
    return response.choices[0].message.content or ""


# ===================================================================
# Response parsing
# ===================================================================
def parse_element_response(raw_response: str, element_key: str) -> dict:
    """
    Parse the LLM JSON response for a single element.
    Returns {"present": bool, "value": str, "raw_response": str}.
    """
    raw_response = raw_response.strip()

    # Try to extract JSON from the response
    json_match = re.search(r"\{[^{}]+\}", raw_response, re.DOTALL)
    if json_match:
        try:
            parsed = json.loads(json_match.group())
            # Determine presence from parsed values
            values = [str(v).strip() for v in parsed.values() if v is not None]
            not_reported_phrases = {"not reported", "not stated", "not mentioned", "n/a", "none", "absent"}
            present = any(
                v.lower() not in not_reported_phrases and v.lower() != "false"
                for v in values
            )
            value_str = "; ".join(v for v in values if v)
            return {"present": present, "value": value_str, "raw_response": raw_response}
        except json.JSONDecodeError:
            pass

    # Fallback: check for keywords
    lower = raw_response.lower()
    if "not reported" in lower or "not stated" in lower or "not mentioned" in lower:
        return {"present": False, "value": "Not reported", "raw_response": raw_response}
    if "unclear" in lower or "illegible" in lower:
        return {"present": True, "value": "Unclear (illegible)", "raw_response": raw_response}

    # If we got text but couldn't parse JSON, assume present
    return {"present": bool(raw_response), "value": raw_response[:200], "raw_response": raw_response}


# ===================================================================
# Annotation encoding (to match human validation format)
# ===================================================================
def encode_annotation(source_present: Optional[int], llm_present: bool) -> object:
    """
    Encode the LLM annotation to match the human validation encoding:
      source=1, llm_present=True  → 1 (TP)
      source=1, llm_present=False → 2 (FN / miss)
      source=0, llm_present=True  → 3 (FP / fabrication)
      source=0, llm_present=False → 'N/A' (TN)
    """
    if source_present is None or pd.isna(source_present):
        return None
    source_present = int(source_present)
    if source_present == 1:
        return 1 if llm_present else 2
    else:
        return 3 if llm_present else "N/A"


# ===================================================================
# Metrics computation (per run)
# ===================================================================
def compute_run_metrics(validation_df: pd.DataFrame, annotator_suffix: str) -> dict:
    """
    Compute per-element and aggregate metrics for a single run.
    Uses the same TP/FN/FP/TN encoding as human_judge_classification_metrics.py.
    """
    results = {}
    total_tp = total_fp = total_fn = total_tn = 0

    for elem in ELEMENTS:
        source_col = elem["source_col"]
        annot_col = f"{elem['key']}_status_{annotator_suffix}"

        if annot_col not in validation_df.columns:
            continue

        mask = validation_df[[source_col, annot_col]].notnull().all(axis=1)
        d = validation_df.loc[mask]

        tp = ((d[source_col] == 1) & (d[annot_col] == 1)).sum()
        fn = ((d[source_col] == 1) & (d[annot_col] == 2)).sum()
        fp = ((d[source_col] == 0) & (d[annot_col] == 3)).sum()
        tn = ((d[source_col] == 0) & (d[annot_col] == "N/A")).sum()

        total = tp + fn + fp + tn
        accuracy = (tp + tn) / total if total > 0 else None
        precision = tp / (tp + fp) if (tp + fp) > 0 else None
        recall = tp / (tp + fn) if (tp + fn) > 0 else None
        f1 = 2 * precision * recall / (precision + recall) if precision and recall and (precision + recall) > 0 else None
        fabrication_rate = fp / total if total > 0 else None

        results[elem["display"]] = {
            "TP": int(tp), "FP": int(fp), "FN": int(fn), "TN": int(tn),
            "total_evaluated": int(total),
            "accuracy": round(accuracy, 4) if accuracy is not None else None,
            "precision": round(precision, 4) if precision is not None else None,
            "recall": round(recall, 4) if recall is not None else None,
            "f1": round(f1, 4) if f1 is not None else None,
            "fabrication_rate": round(fabrication_rate, 4) if fabrication_rate is not None else None,
        }

        total_tp += tp
        total_fp += fp
        total_fn += fn
        total_tn += tn

    grand_total = total_tp + total_fp + total_fn + total_tn
    results["__aggregate__"] = {
        "total_correct": int(total_tp + total_tn),
        "total_fabrications": int(total_fp),
        "total_misses": int(total_fn),
        "total_evaluated": int(grand_total),
        "overall_accuracy": round((total_tp + total_tn) / grand_total, 4) if grand_total > 0 else None,
        "overall_fabrication_rate": round(total_fp / grand_total, 4) if grand_total > 0 else None,
    }

    return results


# ===================================================================
# DeepEval integration
# ===================================================================
def run_deepeval_metrics(source_text: str, llm_output: str, element_name: str) -> dict:
    """
    Run DeepEval hallucination and faithfulness metrics on a single extraction.
    Returns metric scores or None if deepeval is not installed.
    """
    try:
        from deepeval.metrics import HallucinationMetric, FaithfulnessMetric
        from deepeval.test_case import LLMTestCase

        test_case = LLMTestCase(
            input=f"Extract {element_name} from the source documents.",
            actual_output=llm_output,
            context=[source_text[:8000]],  # Truncate for token limits
        )

        hallucination = HallucinationMetric(threshold=0.5)
        faithfulness = FaithfulnessMetric(threshold=0.5)

        hallucination.measure(test_case)
        faithfulness.measure(test_case)

        return {
            "hallucination_score": hallucination.score,
            "hallucination_reason": hallucination.reason,
            "faithfulness_score": faithfulness.score,
            "faithfulness_reason": faithfulness.reason,
        }
    except ImportError:
        return None
    except Exception as exc:
        log.warning("DeepEval metric failed for %s: %s", element_name, exc)
        return None


# ===================================================================
# Single run: one model × one prompt approach × all cases
# ===================================================================
def run_single_configuration(
    model_name: str,
    prompt_approach: str,
    element_prompts: dict,
    system_prompt: str,
    case_dirs: dict,
    validation_df: pd.DataFrame,
    run_deepeval: bool = False,
    dry_run: bool = False,
) -> dict:
    """
    Run extraction for one model × one prompt approach across all matched cases.

    Returns:
        {
            "model": str,
            "prompt_approach": str,
            "timestamp": str,
            "validation_df": pd.DataFrame (with new annotation columns),
            "metrics": dict,
            "per_case_details": list[dict],
            "deepeval_results": list[dict] (if run_deepeval),
        }
    """
    timestamp = datetime.now().isoformat()
    annotator_suffix = f"llm_{model_name.replace('-', '_')}_{prompt_approach}"
    per_case_details = []
    deepeval_results = []

    # Add annotation columns to validation_df
    for elem in ELEMENTS:
        col_name = f"{elem['key']}_status_{annotator_suffix}"
        validation_df[col_name] = None

    for case_key, case_dir in case_dirs.items():
        case_idx = case_key  # row index in validation_df
        if case_idx >= len(validation_df):
            continue

        row = validation_df.iloc[case_idx]
        source_text = get_source_text(case_dir)

        if not source_text:
            log.warning("No source text for case %s", case_dir.name)
            continue

        case_detail = {
            "case_index": case_idx,
            "case_folder": case_dir.name,
            "elements": {},
        }

        for elem in ELEMENTS:
            element_key = elem["key"]
            source_col = elem["source_col"]
            source_val = row.get(source_col)

            if pd.isna(source_val):
                continue

            prompt_text = element_prompts.get(element_key)
            if not prompt_text:
                continue

            # Replace <<<PDF_TEXT>>> or <<<CONTEXT>>> placeholders
            user_prompt = prompt_text.replace("<<<PDF_TEXT>>>", source_text[:12000])
            user_prompt = user_prompt.replace("<<<CONTEXT>>>", source_text[:12000])
            if "<<<" not in prompt_text:
                user_prompt = f"{prompt_text}\n\nSource documents:\n{source_text[:12000]}"

            if dry_run:
                raw_response = '{"value": "DRY_RUN"}'
            else:
                try:
                    raw_response = call_llm(model_name, system_prompt, user_prompt)
                    time.sleep(0.5)  # Rate limiting
                except Exception as exc:
                    log.error("LLM call failed for %s/%s: %s", case_dir.name, element_key, exc)
                    raw_response = ""

            parsed = parse_element_response(raw_response, element_key)
            annotation = encode_annotation(source_val, parsed["present"])

            # Store annotation
            col_name = f"{elem['key']}_status_{annotator_suffix}"
            validation_df.at[case_idx, col_name] = annotation

            case_detail["elements"][elem["display"]] = {
                "source_present": int(source_val) if pd.notna(source_val) else None,
                "llm_present": parsed["present"],
                "llm_value": parsed["value"],
                "annotation": annotation,
                "raw_response_preview": parsed["raw_response"][:300],
            }

            # Optional DeepEval metrics
            if run_deepeval and not dry_run and raw_response:
                de_result = run_deepeval_metrics(source_text, raw_response, elem["display"])
                if de_result:
                    deepeval_results.append({
                        "case_index": case_idx,
                        "element": elem["display"],
                        **de_result,
                    })

        per_case_details.append(case_detail)

    # Compute metrics
    metrics = compute_run_metrics(validation_df, annotator_suffix)

    return {
        "model": model_name,
        "prompt_approach": prompt_approach,
        "timestamp": timestamp,
        "annotator_suffix": annotator_suffix,
        "validation_df": validation_df,
        "metrics": metrics,
        "per_case_details": per_case_details,
        "deepeval_results": deepeval_results,
    }


# ===================================================================
# Build case directory map (indexed by validation row index)
# ===================================================================
def build_indexed_case_map(raw_dir: Path, validation_df: pd.DataFrame) -> dict:
    """Map validation row indices → case directories."""
    # Build folder lookup
    folder_map = {}
    for surgeon_dir in raw_dir.iterdir():
        if not surgeon_dir.is_dir():
            continue
        for case_dir in surgeon_dir.iterdir():
            if not case_dir.is_dir():
                continue
            parts = case_dir.name.split("_")
            if len(parts) >= 2:
                patient_init = parts[1]
                case_type = "invasive" if "invasive" in case_dir.name.lower() else "DCIS"
                folder_map[(surgeon_dir.name, patient_init, case_type)] = case_dir

    # Map row indices
    indexed = {}
    for idx, row in validation_df.iterrows():
        surgeon = row.get("surgeon", "")
        pi = row.get("patient_initials", "")
        tumor = "invasive" if row.get("tumor_invasive_dcis") == 1 else "DCIS"
        surgeon_dir = SURGEON_NAME_TO_DIR.get(surgeon, surgeon)
        key = (surgeon_dir, pi, tumor)
        if key in folder_map:
            indexed[idx] = folder_map[key]

    return indexed


# ===================================================================
# Main
# ===================================================================
def main():
    parser = argparse.ArgumentParser(description="DeepEval Multi-Model Pipeline")
    parser.add_argument("--models", type=str, default=",".join(DEFAULT_MODELS),
                        help="Comma-separated model names")
    parser.add_argument("--approaches", type=str, default=None,
                        help="Comma-separated prompt approaches (default: all)")
    parser.add_argument("--dry-run", action="store_true",
                        help="Skip actual LLM calls; use placeholder responses")
    parser.add_argument("--deepeval", action="store_true",
                        help="Run DeepEval hallucination/faithfulness metrics")
    parser.add_argument("--max-cases", type=int, default=None,
                        help="Limit number of cases to process")
    args = parser.parse_args()

    models = [m.strip() for m in args.models.split(",")]
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    # Load validation data
    log.info("Loading validation data...")
    df_val = pd.read_excel(VALIDATION_RAW)
    log.info("Loaded %d cases", len(df_val))

    # Build case map
    case_dirs = build_indexed_case_map(RAW_DIR, df_val)
    log.info("Matched %d cases to source document folders", len(case_dirs))

    if args.max_cases:
        case_dirs = dict(list(case_dirs.items())[:args.max_cases])
        log.info("Limited to %d cases", len(case_dirs))

    # Load prompt library
    prompt_df = load_prompt_library()

    # Determine approaches to run
    approaches_to_run = []

    # Standard prompt library approaches
    if args.approaches:
        selected = [a.strip() for a in args.approaches.split(",")]
    else:
        selected = PROMPT_APPROACHES

    for approach in selected:
        prompts = get_prompts_for_approach(prompt_df, approach)
        if prompts:
            approaches_to_run.append({
                "name": approach,
                "system_prompt": "You are a clinical data extraction assistant. Extract information from radiology and pathology reports. Output only valid JSON. Do not fabricate information.",
                "element_prompts": prompts,
            })

    # Add BFOP
    approaches_to_run.append({
        "name": "bfop",
        "system_prompt": BFOP_SYSTEM_PROMPT,
        "element_prompts": BFOP_ELEMENT_PROMPTS,
    })

    # Add 2POP
    approaches_to_run.append({
        "name": "2pop",
        "system_prompt": TWO_POP_SYSTEM_PROMPT,
        "element_prompts": TWO_POP_ELEMENT_PROMPTS,
    })

    log.info("Will run %d models × %d approaches = %d configurations",
             len(models), len(approaches_to_run), len(models) * len(approaches_to_run))

    # Run all configurations
    all_run_summaries = []

    for model_name in models:
        for approach_config in approaches_to_run:
            approach_name = approach_config["name"]
            log.info("=" * 60)
            log.info("Running: model=%s, approach=%s", model_name, approach_name)

            result = run_single_configuration(
                model_name=model_name,
                prompt_approach=approach_name,
                element_prompts=approach_config["element_prompts"],
                system_prompt=approach_config["system_prompt"],
                case_dirs=case_dirs,
                validation_df=df_val.copy(),
                run_deepeval=args.deepeval,
                dry_run=args.dry_run,
            )

            # Save run results
            run_id = f"{model_name}_{approach_name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
            run_dir = OUTPUT_DIR / run_id
            run_dir.mkdir(parents=True, exist_ok=True)

            # Save metrics
            metrics_path = run_dir / "metrics.json"
            with open(metrics_path, "w") as f:
                json.dump(result["metrics"], f, indent=2, default=str)

            # Save per-case details
            details_path = run_dir / "per_case_details.json"
            with open(details_path, "w") as f:
                json.dump(result["per_case_details"], f, indent=2, default=str)

            # Save validation sheet (xlsx)
            val_path = run_dir / "validation_sheet.xlsx"
            result["validation_df"].to_excel(val_path, index=False)

            # Save DeepEval results if any
            if result["deepeval_results"]:
                de_path = run_dir / "deepeval_results.json"
                with open(de_path, "w") as f:
                    json.dump(result["deepeval_results"], f, indent=2, default=str)

            # Collect summary
            agg = result["metrics"].get("__aggregate__", {})
            run_summary = {
                "run_id": run_id,
                "model": model_name,
                "prompt_approach": approach_name,
                "timestamp": result["timestamp"],
                "total_correct": agg.get("total_correct"),
                "total_fabrications": agg.get("total_fabrications"),
                "total_misses": agg.get("total_misses"),
                "overall_accuracy": agg.get("overall_accuracy"),
                "overall_fabrication_rate": agg.get("overall_fabrication_rate"),
            }
            all_run_summaries.append(run_summary)

            log.info("  Accuracy: %s  Fabrications: %s",
                     agg.get("overall_accuracy"), agg.get("total_fabrications"))

    # Save master summary
    summary_path = OUTPUT_DIR / "run_summary.json"
    with open(summary_path, "w") as f:
        json.dump(all_run_summaries, f, indent=2, default=str)

    summary_df = pd.DataFrame(all_run_summaries)
    summary_csv = OUTPUT_DIR / "run_summary.csv"
    summary_df.to_csv(summary_csv, index=False)

    log.info("=" * 60)
    log.info("All runs complete. Summary saved to %s", summary_csv)
    log.info("Run details saved under %s", OUTPUT_DIR)


if __name__ == "__main__":
    main()
